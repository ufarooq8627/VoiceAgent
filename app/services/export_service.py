import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import Booking, Candidate, Slot, ConversationMessage, CandidateEvent


class ExportService:
    @classmethod
    def export_job(cls, db: Session, job_id: str) -> openpyxl.Workbook:
        wb = openpyxl.Workbook()

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        header_align = Alignment(horizontal="left", vertical="center")

        def style_header(sheet, headers):
            sheet.append(headers)
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            sheet.row_dimensions[1].height = 24

        def adjust_column_widths(sheet):
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # 1. Schedule Sheet
        ws_schedule = wb.active
        ws_schedule.title = "Schedule"
        style_header(ws_schedule, ["Date", "Time", "Interviewer", "Candidate", "Email", "Mode", "Link"])

        bookings = (
            db.query(Booking)
            .join(Candidate, Booking.candidate_id == Candidate.id)
            .join(Slot, Booking.slot_id == Slot.id)
            .filter(Candidate.job_id == job_id, Booking.status == "confirmed")
            .order_by(Slot.start_at_utc.asc())
            .all()
        )

        for b in bookings:
            slot = b.slot
            cand = b.candidate
            interviewer_name = slot.interviewer.name if slot.interviewer else ""
            date_str = slot.start_at_utc.strftime("%Y-%m-%d")
            time_str = f"{slot.start_at_utc.strftime('%H:%M')} - {slot.end_at_utc.strftime('%H:%M')} UTC"
            ws_schedule.append([
                date_str,
                time_str,
                interviewer_name,
                cand.name,
                cand.email,
                slot.mode,
                slot.location_or_link or "",
            ])
        adjust_column_widths(ws_schedule)

        # 2. Not Scheduled Sheet
        ws_not_scheduled = wb.create_sheet(title="Not Scheduled")
        style_header(ws_not_scheduled, ["Name", "Email", "Status", "Reason"])

        not_booked = (
            db.query(Candidate)
            .filter(Candidate.job_id == job_id, Candidate.status != "BOOKED")
            .order_by(Candidate.created_at.desc())
            .all()
        )

        for c in not_booked:
            # Check latest event reason if any
            event = (
                db.query(CandidateEvent)
                .filter(CandidateEvent.candidate_id == c.id)
                .order_by(CandidateEvent.created_at.desc())
                .first()
            )
            reason = ""
            if event and event.payload and isinstance(event.payload, dict):
                reason = event.payload.get("reason", "")

            ws_not_scheduled.append([
                c.name,
                c.email,
                c.status,
                reason,
            ])
        adjust_column_widths(ws_not_scheduled)

        # 3. Conversation Log Sheet
        ws_conv = wb.create_sheet(title="Conversation Log")
        style_header(ws_conv, ["Candidate", "Direction", "Summary", "Timestamp"])

        messages = (
            db.query(ConversationMessage)
            .join(Candidate, ConversationMessage.candidate_id == Candidate.id)
            .filter(Candidate.job_id == job_id)
            .order_by(ConversationMessage.created_at.asc())
            .all()
        )

        for m in messages:
            cand_name = m.candidate.name if m.candidate else ""
            ts_str = m.created_at.strftime("%Y-%m-%d %H:%M:%S UTC") if m.created_at else ""
            ws_conv.append([
                cand_name,
                m.direction,
                m.summary,
                ts_str,
            ])
        adjust_column_widths(ws_conv)

        return wb
