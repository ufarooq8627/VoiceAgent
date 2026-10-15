import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.orm import Session
from app.models.receptionist import UnansweredEnquiry, KnowledgeItem


class ReceptionistExportService:

    @staticmethod
    def export_enquiries_xlsx(db: Session, business_id: str | None = None) -> io.BytesIO:
        wb = Workbook()
        ws = wb.active
        ws.title = "Unanswered Enquiries"

        headers = [
            "enquiry_id",
            "business_id",
            "business_name",
            "call_id",
            "caller_phone",
            "question",
            "status",
            "owner_answer",
            "created_at",
            "resolved_at",
            "published_to_kb",
        ]
        ws.append(headers)

        # Style Header Row
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Query Enquiries
        query = db.query(UnansweredEnquiry)
        if business_id:
            query = query.filter(UnansweredEnquiry.business_id == business_id)
        enquiries = query.order_by(UnansweredEnquiry.created_at.desc()).all()

        for enq in enquiries:
            # Check if published to KB
            kb_exists = (
                db.query(KnowledgeItem)
                .filter(
                    KnowledgeItem.business_id == enq.business_id,
                    KnowledgeItem.question == enq.question,
                )
                .first()
            )
            published_to_kb = "yes" if kb_exists else "no"

            caller_phone = ""
            if enq.call and enq.call.caller_number:
                caller_phone = enq.call.caller_number

            biz_name = enq.business.name if enq.business else ""

            row_data = [
                enq.id,
                enq.business_id,
                biz_name,
                enq.call_id or "",
                caller_phone,
                enq.question,
                enq.status,
                enq.owner_answer or "",
                enq.created_at.isoformat() if enq.created_at else "",
                enq.resolved_at.isoformat() if enq.resolved_at else "",
                published_to_kb,
            ]
            ws.append(row_data)

            # Apply cell borders
            current_row = ws.max_row
            for col_num in range(1, len(headers) + 1):
                c = ws.cell(row=current_row, column=col_num)
                c.border = thin_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = col[0].column_letter
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def export_enquiries_csv(db: Session, business_id: str | None = None) -> io.StringIO:
        import csv
        output = io.StringIO()
        writer = csv.writer(output)

        headers = [
            "enquiry_id",
            "business_id",
            "business_name",
            "call_id",
            "caller_phone",
            "question",
            "status",
            "owner_answer",
            "created_at",
            "resolved_at",
            "published_to_kb",
        ]
        writer.writerow(headers)

        query = db.query(UnansweredEnquiry)
        if business_id:
            query = query.filter(UnansweredEnquiry.business_id == business_id)
        enquiries = query.order_by(UnansweredEnquiry.created_at.desc()).all()

        for enq in enquiries:
            kb_exists = (
                db.query(KnowledgeItem)
                .filter(
                    KnowledgeItem.business_id == enq.business_id,
                    KnowledgeItem.question == enq.question,
                )
                .first()
            )
            published_to_kb = "yes" if kb_exists else "no"
            caller_phone = enq.call.caller_number if (enq.call and enq.call.caller_number) else ""
            biz_name = enq.business.name if enq.business else ""

            row_data = [
                enq.id,
                enq.business_id,
                biz_name,
                enq.call_id or "",
                caller_phone,
                enq.question,
                enq.status,
                enq.owner_answer or "",
                enq.created_at.isoformat() if enq.created_at else "",
                enq.resolved_at.isoformat() if enq.resolved_at else "",
                published_to_kb,
            ]
            writer.writerow(row_data)

        output.seek(0)
        return output

